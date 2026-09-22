from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt


ROOT = Path(r"G:\stm32_Project\test_projects\PowerPico")
SOURCE = ROOT / "DOC" / "PowerPico_复刻项目_执行规划与进度清单_修订版_升级信息补充.docx"
OUTPUT = ROOT / "tmp" / "PowerPico_进度清单_STEP5更新_待核验.docx"
MARKER = "17.1 Step 5 三档自动换挡（用户硬件实测通过）"


def add_title(doc: Document, text: str, size: float = 14.0) -> None:
    """追加一个与原文风格接近的粗体标题。"""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)


def add_body(doc: Document, text: str) -> None:
    """追加正文，沿用文档默认正文样式。"""
    doc.add_paragraph(text)


def main() -> None:
    """在副本末尾追加STEP5实测记录，原文件留到核验通过后再替换。"""
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    if any(MARKER in paragraph.text for paragraph in doc.paragraphs):
        raise RuntimeError("目标文档已经包含STEP5实测记录，停止重复写入")

    # 新增独立页面，避免破坏原有计划正文和资料链接的排版。
    doc.add_page_break()
    add_title(doc, "17. 阶段进度更新")
    add_title(doc, MARKER, 12.0)

    add_body(
        doc,
        "阶段结论：2026-09-19，用户完成固件烧录和实板负载测试。STEP5核心自动换挡功能验收通过，可以进入Step 6。",
    )
    add_body(
        doc,
        "1. OUT空载：量程最终稳定在LOW。照片中电流约为+43.95μA和-25.20μA，状态为HOLD，升档/降档计数均归零；最近一次记录为M>L、DN。",
    )
    add_body(
        doc,
        "2. 接入ESP32：量程由LOW切换并最终稳定在HIGH。照片中电流约为152.344～155.273mA，状态为HOLD；最近一次记录为M>H、RAIL，升档连续确认门限为4。",
    )
    add_body(
        doc,
        "3. 断开ESP32：量程最终回到LOW。照片中电流约为+31.05～+64.45μA，状态为HOLD；最近一次记录为M>L、DN，降档连续确认门限为10。",
    )
    add_body(
        doc,
        "4. 运行状态：各阶段均显示E0、Q0，照片未见ADC错误或数据块丢弃。S/C分别出现2/10、6/30、8/40，与每次换挡跳过1块、冷却5块的设计关系一致。",
    )

    add_title(doc, "17.2 验收边界与待复测项", 12.0)
    add_body(
        doc,
        "1. 本次通过只确认自动换挡、连续确认、切换后跳过/冷却和数据处理链路能够工作，不代表绝对电流精度、线性、零点和所有临界阈值已经验收。",
    )
    add_body(
        doc,
        "2. LOW档空载/断载读数约为-25.20～+64.45μA，仍存在零点偏移和噪声，留到Step 10使用已知负载统一校准。ESP32约152～155mA只能作为功能性负载，不能替代标准负载精度测试。",
    )
    add_body(
        doc,
        "3. 切换编号从#2到#6之间比单次LOW→HIGH多记录2次切换；仅凭现有照片不能区分重复插拔/测试操作，还是ESP32启动瞬态引起的额外切换。该项不阻塞STEP5核心功能通过，但后续应使用可控负载或连续日志复测临界点稳定性。",
    )

    add_title(doc, "17.3 当前下一步：Step 6 · RTOS任务框架", 12.0)
    add_body(
        doc,
        "在保留STEP5可观察数据的基础上，按官方工程结构逐步拆分ADC处理、UI刷新、按键和通信任务。临时测试代码继续遵循“先简易、能跑起来、再优化”的原则；后续新增或修改代码均补充中文关键注释。",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
